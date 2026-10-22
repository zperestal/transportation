import ast
import re
from datetime import datetime
from pathlib import Path

from docxtpl import DocxTemplate
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parent.parent
PARAMETERS_PATH = PROJECT_DIR / "config" / "parameters.txt"
DATA_DIR = PROJECT_DIR / "data"
TEMPLATE_DIR = PROJECT_DIR / "templates"
OUTPUT_DIR = PROJECT_DIR / "output"


def load_parameters():
	try:
		tree = ast.parse(PARAMETERS_PATH.read_text(encoding="utf-8"), filename=str(PARAMETERS_PATH))
	except (OSError, SyntaxError) as error:
		raise ValueError(f"Could not read parameters from {PARAMETERS_PATH}: {error}") from error

	parameters = {}
	for statement in tree.body:
		if not isinstance(statement, ast.Assign) or len(statement.targets) != 1:
			continue
		target = statement.targets[0]
		if not isinstance(target, ast.Name):
			continue
		try:
			parameters[target.id] = ast.literal_eval(statement.value)
		except (TypeError, ValueError) as error:
			raise ValueError(f"Parameter {target.id!r} must have a literal value.") from error

	required = ("src_name_excel", "tmpl_name_word", "sheet_to_read", "line_to_read")
	missing = [name for name in required if name not in parameters]
	if missing:
		raise ValueError(f"Missing parameter(s) in {PARAMETERS_PATH}: {', '.join(missing)}")
	if not isinstance(parameters["line_to_read"], int) or isinstance(parameters["line_to_read"], bool):
		raise ValueError("line_to_read must be an integer row number.")
	if parameters["line_to_read"] < 2:
		raise ValueError("line_to_read must be 2 or greater.")
	if not isinstance(parameters["sheet_to_read"], str) or not parameters["sheet_to_read"].strip():
		raise ValueError("sheet_to_read must be a non-empty worksheet name.")
	return parameters


def normalize_field_name(field_name):
	name = re.sub(r"[^a-zA-Z0-9_]+", "_", field_name.strip().lower()).strip("_")
	if name and name[0].isdigit():
		name = f"field_{name}"
	return name


def read_data_row(worksheet, line_to_read):
	if line_to_read > worksheet.max_row:
		raise ValueError(
			f"Row {line_to_read} does not exist in worksheet {worksheet.title!r}."
		)

	header_row_number = None
	headers = None
	for row_number in range(1, line_to_read):
		candidate = [cell.value for cell in worksheet[row_number]]
		if any(value is not None and str(value).strip() for value in candidate):
			header_row_number = row_number
			headers = candidate
			break

	if headers is None:
		raise ValueError(
			f"No field-name row was found before row {line_to_read} in worksheet {worksheet.title!r}."
		)
	if line_to_read <= header_row_number:
		raise ValueError(
			f"line_to_read must identify a data row after the field names in {worksheet.title!r}."
		)

	values = [cell.value for cell in worksheet[line_to_read]]
	column_values = {}
	field_values = {}
	for column_number, value in enumerate(values, start=1):
		value = "" if value is None else value
		column_values[get_column_letter(column_number)] = value
		header = headers[column_number - 1] if column_number <= len(headers) else None
		if header is None or not str(header).strip():
			continue
		field_name = str(header).strip()
		if field_name in field_values:
			raise ValueError(f"Duplicate field name in worksheet header: {field_name!r}")
		field_values[field_name] = value

	return column_values, field_values


def build_agent_context(workbook, load_fields):
	if "Agents" not in workbook.sheetnames:
		raise ValueError("The template requires an 'Agents' worksheet, but it is missing.")

	worksheet = workbook["Agents"]
	if worksheet.max_row < 2:
		raise ValueError("The 'Agents' worksheet has no agent records.")

	headers = [cell.value for cell in worksheet[1]]
	header_columns = {
		normalize_field_name(str(header)): index
		for index, header in enumerate(headers)
		if header is not None and str(header).strip()
	}
	name_column = header_columns.get("name")
	city_column = header_columns.get("city")
	if name_column is None or city_column is None:
		raise ValueError("The 'Agents' worksheet must have 'Name' and 'City' columns.")

	load_name = load_fields.get("saler_name") or load_fields.get("seller_name")
	load_city = load_fields.get("saler_city") or load_fields.get("seller_city")
	if not load_name or not load_city:
		raise ValueError("The selected load row must contain seller name and city fields.")

	for row_number in range(2, worksheet.max_row + 1):
		values = [cell.value for cell in worksheet[row_number]]
		name = values[name_column]
		city = values[city_column]
		if (
			str(name or "").strip().casefold() == str(load_name).strip().casefold()
			and str(city or "").strip().casefold() == str(load_city).strip().casefold()
		):
			return {
				get_column_letter(column_number): "" if value is None else value
				for column_number, value in enumerate(values, start=1)
			}

	raise ValueError(
			f"No agent matches seller {load_name!r} in {load_city!r}."
	)


def build_template_context(column_values, field_values, agent_values):
	context = {
		"loads": column_values,
		"Agents": agent_values,
		"row": field_values,
	}
	for field_name, value in field_values.items():
		placeholder_name = normalize_field_name(field_name)
		if placeholder_name:
			context.setdefault(placeholder_name, value)
	return context


def main():
	parameters = load_parameters()
	excel_path = DATA_DIR / parameters["src_name_excel"]
	template_path = TEMPLATE_DIR / parameters["tmpl_name_word"]
	if not excel_path.is_file():
		raise FileNotFoundError(f"Excel file not found: {excel_path}")
	if not template_path.is_file():
		raise FileNotFoundError(f"Word template not found: {template_path}")

	line_to_read = parameters["line_to_read"]
	workbook = load_workbook(excel_path, data_only=True)
	try:
		sheet_name = parameters["sheet_to_read"]
		if sheet_name not in workbook.sheetnames:
			raise ValueError(
				f"Worksheet {sheet_name!r} not found in {excel_path.name}. "
			f"Available worksheets: {', '.join(workbook.sheetnames)}"
			)
		column_values, field_values = read_data_row(workbook[sheet_name], line_to_read)
		agent_values = build_agent_context(workbook, field_values)
	finally:
		workbook.close()

	context = build_template_context(column_values, field_values, agent_values)
	identifier = column_values.get("B")
	if identifier is None or not str(identifier).strip():
		raise ValueError(f"Cell B{line_to_read} is empty in worksheet {sheet_name!r}.")
	identifier = re.sub(r'[\\/*?:"<>|]', "_", str(identifier).strip())
	timestamp = datetime.now().strftime("%Y%m%d_%H%M")
	output_path = OUTPUT_DIR / f"Load_{identifier}_{timestamp}.docx"
	OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

	document = DocxTemplate(template_path)
	document.render(context)
	document.save(output_path)

	print(f"Created Word document: {output_path}")


if __name__ == "__main__":
	try:
		main()
	except (FileNotFoundError, OSError, ValueError) as error:
		raise SystemExit(f"Error: {error}") from error
